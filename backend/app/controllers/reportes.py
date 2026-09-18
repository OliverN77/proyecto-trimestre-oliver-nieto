import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.exceptions import AuthorizationError
from app.core.responses import API_RESPONSES
from app.database import get_db
from app.dependencies.auth import get_current_user
from app.models.mesa import Mesa
from app.models.reserva import Reserva
from app.models.reserva_producto import ReservaProducto
from app.models.reserva_servicio import ReservaServicio
from app.models.usuario import Usuario

router = APIRouter(prefix="/api/reportes", tags=["Reportes"], responses=API_RESPONSES)


def _get_daily_reservations(db: Session, fecha: date):
    query = (
        select(Reserva, Mesa.numero_mesa, Usuario)
        .options(
            selectinload(Reserva.productos).selectinload(ReservaProducto.producto),
            selectinload(Reserva.servicios).selectinload(ReservaServicio.servicio),
        )
        .join(Mesa, Mesa.id_mesa == Reserva.id_mesa)
        .join(Usuario, Usuario.id_usuario == Reserva.id_cliente)
        .where(Reserva.fecha_reserva == fecha)
        .order_by(Reserva.hora_inicio)
    )
    return db.execute(query).all()


@router.get("/reservas/diario", summary="Reporte diario de reservas (JSON)")
def reporte_diario_json(
    fecha: date = Query(default_factory=date.today),
    usuario: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if usuario.id_rol not in (1, 2):
        raise AuthorizationError()

    rows = _get_daily_reservations(db, fecha)
    reservas = []
    total_general = 0.0
    for reserva, numero_mesa, cliente in rows:
        reservas.append({
            "id_reserva": reserva.id_reserva,
            "cliente": f"{cliente.nombre} {cliente.apellido}",
            "mesa": numero_mesa,
            "cantidad_personas": reserva.cantidad_personas,
            "hora_inicio": str(reserva.hora_inicio)[:5],
            "hora_fin": str(reserva.hora_fin)[:5],
            "estado": reserva.estado,
            "subtotal": reserva.subtotal,
            "impuestos": reserva.impuestos,
            "total": reserva.total,
            "productos": [f"{rp.cantidad}x {rp.producto.nombre}" for rp in reserva.productos] if reserva.productos else [],
            "servicios": [f"{rs.cantidad}x {rs.servicio.nombre}" for rs in reserva.servicios] if reserva.servicios else [],
        })
        total_general += reserva.total

    return {
        "fecha": str(fecha),
        "total_reservas": len(reservas),
        "total_general": round(total_general, 2),
        "reservas": reservas,
    }


@router.get("/reservas/diario/pdf", summary="Reporte diario de reservas (PDF)")
def reporte_diario_pdf(
    fecha: date = Query(default_factory=date.today),
    usuario: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if usuario.id_rol not in (1, 2):
        raise AuthorizationError()

    rows = _get_daily_reservations(db, fecha)

    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph("Taberna del Faro", styles["Title"]))
    elements.append(Paragraph(f"Reporte Diario de Reservas — {fecha}", styles["Heading2"]))
    elements.append(Spacer(1, 12))

    # Table header
    data = [["#", "Cliente", "Mesa", "Personas", "Horario", "Estado", "Total"]]
    total_general = 0.0
    for reserva, numero_mesa, cliente in rows:
        data.append([
            str(reserva.id_reserva),
            f"{cliente.nombre} {cliente.apellido}",
            str(numero_mesa),
            str(reserva.cantidad_personas),
            f"{str(reserva.hora_inicio)[:5]}–{str(reserva.hora_fin)[:5]}",
            reserva.estado,
            f"${reserva.total:,.2f}",
        ])
        total_general += reserva.total

    data.append(["", "", "", "", "", "TOTAL", f"${total_general:,.2f}"])

    table = Table(data, colWidths=[35, 120, 40, 55, 80, 70, 70])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -2), colors.HexColor("#f6efdd")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e8b23d")),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"Generado el {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC", styles["Normal"]))

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=reporte_reservas_{fecha}.pdf"},
    )


@router.get("/reservas/diario/excel", summary="Reporte diario de reservas (Excel)")
def reporte_diario_excel(
    fecha: date = Query(default_factory=date.today),
    usuario: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if usuario.id_rol not in (1, 2):
        raise AuthorizationError()

    rows = _get_daily_reservations(db, fecha)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = f"Reservas {fecha}"

    # Header styling
    header_fill = PatternFill(start_color="16324f", end_color="16324f", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)

    # Title
    ws.merge_cells("A1:G1")
    ws["A1"] = f"Taberna del Faro — Reporte Diario de Reservas — {fecha}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["# Reserva", "Cliente", "Mesa", "Personas", "Horario", "Estado", "Total"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    total_general = 0.0
    for i, (reserva, numero_mesa, cliente) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=reserva.id_reserva)
        ws.cell(row=i, column=2, value=f"{cliente.nombre} {cliente.apellido}")
        ws.cell(row=i, column=3, value=numero_mesa)
        ws.cell(row=i, column=4, value=reserva.cantidad_personas)
        ws.cell(row=i, column=5, value=f"{str(reserva.hora_inicio)[:5]}–{str(reserva.hora_fin)[:5]}")
        ws.cell(row=i, column=6, value=reserva.estado)
        ws.cell(row=i, column=7, value=reserva.total)
        total_general += reserva.total

    last_row = len(rows) + 4
    ws.cell(row=last_row, column=6, value="TOTAL").font = Font(bold=True)
    ws.cell(row=last_row, column=7, value=total_general).font = Font(bold=True)

    # Column widths
    for col_letter, width in [("A", 12), ("B", 25), ("C", 8), ("D", 10), ("E", 15), ("F", 12), ("G", 12)]:
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=reporte_reservas_{fecha}.xlsx"},
    )
